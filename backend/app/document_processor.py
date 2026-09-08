"""
document_processor.py
----------------------
Responsible for the "document ingestion" hands-on concepts this project
demonstrates:

    1. Text extraction (PDF, Word, PowerPoint)
    2. Document chunking

Design notes:
- Text is extracted UNIT BY UNIT (PDF page / Word "page" / PPT slide) so
  that every chunk can keep track of exactly where it came from. That
  location is what later powers citations like "Source: paper.pdf, Page 5"
  or "Source: deck.pptx, Slide 3".
- Chunking is implemented manually (simple sliding window over
  characters) rather than pulling in a heavy text-splitting library.
  This keeps the dependency list small and makes the chunking logic
  fully transparent/inspectable for a mini-project.
- Word documents don't have a reliable notion of "pages" without actually
  rendering them (page breaks depend on fonts, margins, printer settings,
  etc.). We take a best-effort approach: if the .docx contains explicit
  manual page breaks, we honor them and label citations "Page N". If it
  doesn't, we fall back to labeling citations "Section N", where each
  section is just a sequential slice of the document -- so citations
  stay honest about what they represent instead of pretending to be
  real page numbers.
"""

from dataclasses import dataclass
from typing import List
import os
import uuid

from pypdf import PdfReader
from pypdf.errors import PdfReadError

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from pptx import Presentation

from app.config import CHUNK_SIZE, CHUNK_OVERLAP

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx")


@dataclass
class Chunk:
    """A single chunk of document text plus the metadata needed for citations."""
    chunk_id: str
    text: str
    doc_name: str
    page_number: int
    chunk_index: int  # index of this chunk within its page/slide/section
    unit_label: str = "Page"  # "Page" (PDF / Word with page breaks), "Section" (Word without page breaks), "Slide" (PPTX)

    def metadata(self) -> dict:
        return {
            "doc_name": self.doc_name,
            "page_number": self.page_number,
            "chunk_index": self.chunk_index,
            "unit_label": self.unit_label,
        }


class DocumentProcessingError(Exception):
    """Raised when a document cannot be read or contains no extractable text."""


# Kept as an alias so any existing code/imports referring to the old,
# PDF-only name keep working now that this module handles more formats.
PDFProcessingError = DocumentProcessingError


# ---------------------------------------------------------------------------
# PDF extraction
# ---------------------------------------------------------------------------
def _extract_pdf_pages(file_bytes, doc_name: str) -> List[dict]:
    """Extract text from every page of a PDF. Returns 1-indexed pages."""
    try:
        reader = PdfReader(file_bytes)
    except (PdfReadError, Exception) as exc:  # pypdf can raise various errors
        raise DocumentProcessingError(f"Could not read '{doc_name}': {exc}") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")  # try an empty password before giving up
        except Exception:
            pass
        if reader.is_encrypted:
            raise DocumentProcessingError(f"'{doc_name}' is password-protected and cannot be read.")

    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append({"unit_number": i, "text": text.strip(), "unit_label": "Page"})

    if not any(p["text"] for p in pages):
        raise DocumentProcessingError(
            f"No extractable text found in '{doc_name}'. "
            "It may be a scanned/image-only PDF with no text layer."
        )

    return pages


# ---------------------------------------------------------------------------
# DOCX extraction
# ---------------------------------------------------------------------------
def _paragraph_has_page_break(paragraph) -> bool:
    """True if a manual page break run is present in this paragraph."""
    for run in paragraph.runs:
        for br in run._element.findall(qn("w:br")):
            if br.get(qn("w:type")) == "page":
                return True
    return False


def _extract_docx_pages(file_bytes, doc_name: str) -> List[dict]:
    """
    Extract text from a Word document.

    Walks paragraphs and tables in document order, splitting into "pages"
    on manual page breaks when present. If the document has no manual
    page breaks (the common case for docs written without printing in
    mind), the whole document is split into fixed-size "sections" instead,
    so citations don't imply a page number we can't actually guarantee.
    """
    try:
        document = DocxDocument(file_bytes)
    except Exception as exc:
        raise DocumentProcessingError(f"Could not read '{doc_name}': {exc}") from exc

    parts: List[str] = []
    page_breaks_at: List[int] = []  # index into `parts` where a page break occurred

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            para = Paragraph(child, document)
            if para.text.strip():
                parts.append(para.text.strip())
            if _paragraph_has_page_break(para):
                page_breaks_at.append(len(parts))
        elif child.tag == qn("w:tbl"):
            from docx.table import Table
            table = Table(child, document)
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

    full_text = "\n".join(parts).strip()
    if not full_text:
        raise DocumentProcessingError(
            f"No extractable text found in '{doc_name}'. It may be empty or contain only images."
        )

    if page_breaks_at:
        # Honor real page breaks: rebuild page boundaries from the paragraph indices.
        pages = []
        start = 0
        for boundary in page_breaks_at + [len(parts)]:
            segment = "\n".join(parts[start:boundary]).strip()
            if segment:
                pages.append(segment)
            start = boundary
        return [
            {"unit_number": i, "text": text, "unit_label": "Page"}
            for i, text in enumerate(pages, start=1)
            if text
        ]

    # No manual page breaks found: fall back to fixed-size pseudo-sections
    # so long documents still get multiple, manageable citation units.
    section_size = CHUNK_SIZE * 3
    normalized = " ".join(full_text.split())
    sections = [normalized[i:i + section_size].strip() for i in range(0, len(normalized), section_size)]
    return [
        {"unit_number": i, "text": text, "unit_label": "Section"}
        for i, text in enumerate(sections, start=1)
        if text
    ]


# ---------------------------------------------------------------------------
# PPTX extraction
# ---------------------------------------------------------------------------
def _extract_pptx_pages(file_bytes, doc_name: str) -> List[dict]:
    """Extract text from every slide of a PowerPoint deck (1-indexed)."""
    try:
        presentation = Presentation(file_bytes)
    except Exception as exc:
        raise DocumentProcessingError(f"Could not read '{doc_name}': {exc}") from exc

    pages = []
    for i, slide in enumerate(presentation.slides, start=1):
        parts = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        parts.append(row_text)

        # Speaker notes often carry useful context (e.g. talking points).
        if slide.has_notes_slide:
            notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes_text:
                parts.append(f"[Speaker notes] {notes_text}")

        pages.append({"unit_number": i, "text": "\n".join(parts).strip(), "unit_label": "Slide"})

    if not any(p["text"] for p in pages):
        raise DocumentProcessingError(
            f"No extractable text found in '{doc_name}'. Slides may contain only images."
        )

    return pages


# ---------------------------------------------------------------------------
# Shared chunking + pipeline
# ---------------------------------------------------------------------------
def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """
    Split a block of text into overlapping character-based chunks.

    A simple sliding window: each chunk is `chunk_size` characters, and
    consecutive chunks overlap by `overlap` characters so that context
    near chunk boundaries isn't lost. We snap chunk boundaries to the
    nearest whitespace where possible to avoid slicing words in half.
    """
    text = " ".join(text.split())  # normalize whitespace
    if not text:
        return []

    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    text_len = len(text)

    while start < text_len:
        end = min(start + chunk_size, text_len)

        # Try to end the chunk at a space rather than mid-word.
        if end < text_len:
            last_space = text.rfind(" ", start, end)
            if last_space > start:
                end = last_space

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= text_len:
            break

        # Move the window forward, stepping back by `overlap` for context continuity.
        start = max(end - overlap, start + 1)

    return chunks


def process_document(file_bytes, doc_name: str) -> List[Chunk]:
    """
    Full pipeline for one document: extract units (pages/slides/sections)
    -> chunk each unit -> attach metadata.

    Supports .pdf, .docx, and .pptx (dispatched by file extension).
    Returns a flat list of Chunk objects ready for embedding + storage.
    """
    ext = os.path.splitext(doc_name)[1].lower()

    if ext == ".pdf":
        units = _extract_pdf_pages(file_bytes, doc_name)
    elif ext == ".docx":
        units = _extract_docx_pages(file_bytes, doc_name)
    elif ext == ".pptx":
        units = _extract_pptx_pages(file_bytes, doc_name)
    else:
        raise DocumentProcessingError(
            f"Unsupported file type '{ext}' for '{doc_name}'. "
            f"Supported types: {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    all_chunks: List[Chunk] = []
    for unit in units:
        unit_chunks = chunk_text(unit["text"])
        for idx, chunk_str in enumerate(unit_chunks):
            all_chunks.append(
                Chunk(
                    chunk_id=str(uuid.uuid4()),
                    text=chunk_str,
                    doc_name=doc_name,
                    page_number=unit["unit_number"],
                    chunk_index=idx,
                    unit_label=unit["unit_label"],
                )
            )

    if not all_chunks:
        raise DocumentProcessingError(f"'{doc_name}' produced no usable chunks after processing.")

    return all_chunks


# ---------------------------------------------------------------------------
# Backward-compatible aliases (in case other code imports the old names)
# ---------------------------------------------------------------------------
def process_pdf(file_bytes, doc_name: str) -> List[Chunk]:
    """Deprecated alias -- use process_document(). Kept for backward compatibility."""
    return process_document(file_bytes, doc_name)
